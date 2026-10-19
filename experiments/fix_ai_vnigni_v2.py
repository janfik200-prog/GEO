# -*- coding: utf-8 -*-
"""Правка колоды «ИИ во ВНИГНИ» — версия коллег от 28.09.2026 (переоформленная).

ЗАЧЕМ ОТДЕЛЬНЫЙ СКРИПТ, А НЕ ПЕРЕСБОРКА. Присланный файл переоформлен вне
репозитория: холст 10 x 5.625", светлый институтский стиль, синий акцент,
верхний и нижний колонтитулы. Сборочный скрипт
``experiments/build_ai_vnigni_presentation.py`` строит совсем другую вёрстку
(13.333 x 7.5", тёмная шапка), поэтому пересборка уничтожила бы оформление.
Здесь исходник открывается как есть и правится точечно; сам исходник в
загрузках не меняется, результат пишется в ``docs II presentation/``.
Скрипт идемпотентен: он всегда читает SRC и целиком перерисовывает содержимое
слайдов 2 и 3, не накапливая правок.

ЧТО ПРАВИТСЯ.
Слайд 2 «Зачем нужен новый метод» — содержание углублено:
  * дорожка эталона: шесть факторов формулы (расстояния до объектов
    геологической основы), нормировка и свёртка нечёткой мерой Плюты;
  * дорожка нового метода: источники съёмок, расчёт признаков, отбор
    (105 -> 58 -> 37), модель и сглаживание карты;
  * добавлена третья дорожка — проверка результата на полосах с разрывом
    15 км с поправкой на региональный наклон.
Слайд 3 «Как построена работа» — схема перерисована под карту организации
  из 11 блоков с ветками 8A-8E и контуром возврата, сгруппирована по стадиям
  (подготовка / данные и знания / эксперименты / верификация и итог).
Слайд 4 — заголовок «Где мы сейчас» -> «Текущие результаты работ»,
  метка «ЭТАП» -> «РЕЗУЛЬТАТЫ». Содержание и карта не тронуты.
Слайд 5 — заголовок «Что дальше» -> «План дальнейших работ». Содержание
  не тронуто.

ИСТОЧНИКИ ЧИСЕЛ И ФОРМУЛИРОВОК (всё проверено по файлам репозитория):
  * шесть факторов формулы и их состав — ``experiments/criterial_reconstruct_multisheet.py``,
    словарь ``ROLE_LAYER`` (tect1 NW_razl, tect2 NE_razl, facies fazii_dist,
    paleo paleo_dist, struct kory_dist, magm dayki_all_dist) и
    ``src/config.py::TAXONOMY_WEIGHTS`` (те же шесть ключей с весами);
    «разломные зоны» (не линии) — memory ``project_fault_layers_are_zones``;
  * мера Плюты — ``src/features.py``, ``taxonomy_weighted_distance``
    («взвешенное L1-расстояние ГИС Интегро (мера Плюты)») и строка
    «нормировка Плюты: меньше = ближе к эталону»;
  * 105 рассчитанных признаков — ``outputs/metrics/honest_feature_signal.csv``
    (сумма размеров групп в столбце «набор»: s2 17 + ter 6 + l8 15 + s2raw 7 +
    relief2 3 + dist 2 + s1 8 + lin 6 + psr 10 + opt 2 + ast 2 + pf 12 + gm 15
    = 105; то же число в докстринге ``experiments/honest_feature_selection.py``);
  * 58 и 37 — ``data/processed/honest_feature_set.json``, поля «признаки» (58)
    и «признаки_перенос» (37);
  * состав источников съёмок — группы того же набора: ter (рельеф), s2/s2raw
    (Sentinel-2), s1 (радар Sentinel-1), l8 (Landsat-8), dist (гидросеть),
    lin (линеаменты); ASTER посчитан отдельно
    (``data/processed/ast_features.parquet``, 26 столбцов группы ast в
    ``dataset_v5_rebuilt.parquet``);
  * гистограммный бустинг — ``experiments/forecast_wide.py``
    (``HistBackgroundEnsemble``, «гистограммный бустинг»);
  * направленное ядро 5 x 1.25 км — ``experiments/plot_ai_vnigni_maps.py``
    (докстринг «ядро 5 x 1.25 км»), северо-западное простирание —
    ``experiments/branch3_shapes.py`` («ядро вдоль северо-западного
    простирания (315 градусов)»);
  * полосы с разрывом 15 км — ``experiments/honest_feature_signal.py``
    (``STRIP_BUFFER_M``), поправка на наклон — ``совпадение_без_тренда``
    в ``experiments/branch1_post_smoothing_sweep.py`` (из обеих карт вычитается
    координатный тренд, затем сравниваются верхние 10% площади);
  * 11 блоков и ветки 8A-8E — ``experiments/build_tech_scheme_presentation.py``
    (списки ``NODES``, ``EDGES``, ``BRANCHES``, контур возврата 4 -> 3).

Запуск из корня: ``python -X utf8 -m experiments.fix_ai_vnigni_v2``
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SRC = Path(r"C:\Users\gstat\Downloads\Telegram Desktop") / \
    "Презентация_ИИ_во_ВНИГНИ_28_09_2026_2 (2).pptx"
OUT = ROOT / "docs II presentation" / "Презентация-ИИ-во-ВНИГНИ-06.10.2026.pptx"

# Палитра присланного файла (снята со слайдов 2-3 исходника).
ACCENT = RGBColor(0x25, 0x63, 0xEB)   # синий акцент
ACCENT_D = RGBColor(0x1D, 0x4E, 0xD8)  # текст на бледно-голубом
PALE_BLUE = RGBColor(0xDB, 0xEA, 0xFE)  # заливка «в работе»
ON_BLUE = RGBColor(0xDB, 0xEA, 0xFE)   # подпись на синем
DARK = RGBColor(0x0F, 0x17, 0x2A)     # основной текст
DARK_BOX = RGBColor(0x1E, 0x29, 0x3B)  # тёмный блок
ON_DARK = RGBColor(0xCB, 0xD5, 0xE1)   # подпись на тёмном
GREY = RGBColor(0x47, 0x55, 0x69)
GREY_L = RGBColor(0x94, 0xA3, 0xB8)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
BORDER_G = RGBColor(0xCB, 0xD5, 0xE1)
PALE = RGBColor(0xF8, 0xFA, 0xFC)     # заливка «впереди» / блок-источник
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"


# --- примитивы в стиле присланного файла -----------------------------------
def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75):
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.text_frame.word_wrap = True
    return sh


def line(sl, x1, y1, x2, y2, color=GREY_L, lw=1.0, arrow=True):
    """Прямая линия (prstGeom line) — тот же приём, что в исходнике."""
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(min(x1, x2)),
                             Inches(min(y1, y2)),
                             Inches(abs(x2 - x1)), Inches(abs(y2 - y1)))
    sh.shadow.inherit = False
    prst = sh._element.spPr.find(qn("a:prstGeom"))
    prst.set("prst", "line")
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(lw)
    if arrow:
        ln = sh._element.spPr.find(qn("a:ln"))
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
        ln.append(tail)
    # для линий справа налево / снизу вверх нужен разворот рамки
    if x2 < x1 or y2 < y1:
        xfrm = sh._element.spPr.find(qn("a:xfrm"))
        if x2 < x1:
            xfrm.set("flipH", "1")
        if y2 < y1:
            xfrm.set("flipV", "1")
    return sh


def textbox(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    sh = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def run(par, text, size, color, bold=False, italic=False, spacing=None):
    r = par.add_run()
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if spacing:
        r._r.get_or_add_rPr().set("spc", str(int(spacing)))
    return r


def caps(sl, x, y, w, text, color, size=8.5):
    """Надпись-рубрика: заглавные, разреженные, жирные — как в исходнике."""
    tf = textbox(sl, x, y, w, 0.16)
    run(tf.paragraphs[0], text, size, color, bold=True, spacing=100)
    return tf


def node(sl, x, y, w, h, title, sub=None, fill=ACCENT, border=None,
         tcolor=WHITE, scolor=ON_BLUE, num=None, ncolor=None,
         align=PP_ALIGN.CENTER, tsize=8.5, ssize=7.5, pad=0.05):
    """Блок схемы: прямоугольник + подпись из заголовка и курсивного пояснения."""
    rect(sl, x, y, w, h, fill, border)
    tf = textbox(sl, x + pad, y, w - 2 * pad, h, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = align
    if num:
        run(p, num + "  ", tsize, ncolor or scolor, bold=True)
    run(p, title, tsize, tcolor, bold=True)
    if sub:
        ps = tf.add_paragraph()
        ps.alignment = align
        run(ps, sub, ssize, scolor, italic=True)
    return tf


# --- служебное --------------------------------------------------------------
def drop_shapes(slide, ids):
    """Удалить фигуры по shape_id (чтобы не трогать шапку, подвал и врезку)."""
    for sh in list(slide.shapes):
        if sh.shape_id in ids:
            sh._element.getparent().remove(sh._element)


def retext(slide, old, new):
    """Заменить текст, не теряя оформления: правим run.text, а не text_frame."""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text == old:
                    r.text = new
                    return True
    raise RuntimeError(f"не найден текст {old!r}")


# ===========================================================================
def slide2(sl):
    """Зачем нужен новый метод: эталон, новый метод, проверка."""
    drop_shapes(sl, set(range(11, 34)))   # всё содержимое между подзаголовком и врезкой

    def summary(yy, text):
        """Подпись-итог дорожки — справа в одной строке с её рубрикой."""
        tf = textbox(sl, 4.50, yy, 5.0, 0.21)
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        run(tf.paragraphs[0], text, 8.5, GREY, italic=True)

    # --- дорожка 1: эталон -------------------------------------------------
    caps(sl, 0.50, 1.80, 4.0, "КРИТЕРИАЛЬНЫЙ МЕТОД ВНИГНИ · ЭТАЛОН", DARK_BOX)
    summary(1.80, "Опирается на подробную геологическую основу листа")
    y, h = 1.97, 0.56
    node(sl, 0.50, y, 2.60, h, "Геологическая основа листа",
         "разломные зоны двух систем, фации, палеоструктуры, "
         "коры выветривания, дайки",
         fill=PALE, border=BORDER, tcolor=DARK, scolor=GREY,
         align=PP_ALIGN.LEFT, tsize=9.5, pad=0.12)
    line(sl, 3.14, y + h / 2, 3.41, y + h / 2)
    node(sl, 3.45, y, 3.45, h, "Критериальная формула",
         "шесть факторов — расстояния до объектов основы; значения "
         "нормируются и сворачиваются нечёткой мерой Плюты с весами",
         fill=DARK_BOX, tcolor=WHITE, scolor=ON_DARK, tsize=9.5)
    line(sl, 6.94, y + h / 2, 7.21, y + h / 2)
    node(sl, 7.25, y, 2.25, h, "Карта перспективности",
         fill=DARK_BOX, tcolor=WHITE, tsize=9.5)

    # --- дорожка 2: новый метод -------------------------------------------
    caps(sl, 0.50, 2.68, 4.0, "НОВЫЙ МЕТОД · ОТКРЫТЫЕ ДАННЫЕ", ACCENT)
    summary(2.68, "Открытые данные есть на любую территорию")
    y, h = 2.85, 0.70
    node(sl, 0.50, y, 2.60, h, "Рельеф и спутниковая съёмка",
         "рельеф, Sentinel-2, Sentinel-1 (радар), Landsat-8, ASTER, гидросеть",
         fill=PALE, border=BORDER, tcolor=DARK, scolor=GREY,
         align=PP_ALIGN.LEFT, tsize=9.5, pad=0.12)
    line(sl, 3.14, y + h / 2, 3.41, y + h / 2)
    node(sl, 3.45, y, 1.75, h, "Расчёт признаков",
         "спектральные индексы, минеральные отношения, морфометрия "
         "рельефа, текстуры, линеаменты", tsize=9.5)
    line(sl, 5.24, y + h / 2, 5.51, y + h / 2)
    node(sl, 5.55, y, 1.75, h, "Отбор признаков",
         "из 105 рассчитанных отобраны 58 на листе; 37 переносятся "
         "на соседние листы", tsize=9.5)
    line(sl, 7.34, y + h / 2, 7.61, y + h / 2)
    node(sl, 7.65, y, 1.85, h, "Модель и карта",
         "гистограммный бустинг; карта сглаживается направленным "
         "ядром 5 × 1.25 км", tsize=9.5)

    # --- дорожка 3: проверка ----------------------------------------------
    caps(sl, 0.50, 3.70, 4.0, "ПРОВЕРКА РЕЗУЛЬТАТА", ACCENT)
    summary(3.70, "Проверка идёт на участках, которых модель не видела")
    y, h = 3.87, 0.50
    node(sl, 0.50, y, 2.60, h, "Полосы с разрывом 15 км",
         "модель учится на одних полосах листа, проверяется на других",
         fill=PALE, border=BORDER, tcolor=DARK, scolor=GREY,
         align=PP_ALIGN.LEFT, tsize=9.5, pad=0.12)
    line(sl, 3.14, y + h / 2, 3.41, y + h / 2)
    node(sl, 3.45, y, 3.45, h, "Сопоставление с картой ВНИГНИ",
         "сравниваются верхние 10% площади обеих карт",
         fill=PALE_BLUE, border=ACCENT, tcolor=ACCENT_D, scolor=ACCENT_D,
         tsize=9.5)
    line(sl, 6.94, y + h / 2, 7.21, y + h / 2)
    node(sl, 7.25, y, 2.25, h, "Поправка на наклон",
         "из обеих карт вычитается общий тренд по координатам",
         fill=PALE_BLUE, border=ACCENT, tcolor=ACCENT_D, scolor=ACCENT_D,
         tsize=9.5)


# ===========================================================================
# Статусы блоков: «готово» — синяя заливка, «в работе» — бледно-голубая,
# «впереди» — светло-серая (легенда слайда остаётся из исходника).
ST = {
    "готово":  dict(fill=ACCENT, border=ACCENT, tcolor=WHITE, scolor=ON_BLUE,
                    ncolor=ON_BLUE),
    "в работе": dict(fill=PALE_BLUE, border=ACCENT, tcolor=ACCENT_D,
                     scolor=ACCENT_D, ncolor=ACCENT),
    "впереди": dict(fill=PALE, border=BORDER_G, tcolor=GREY, scolor=GREY,
                    ncolor=GREY_L),
}


def slide3(sl):
    """Как построена работа: 11 блоков организации, сгруппированных по стадиям."""
    drop_shapes(sl, set(range(17, 61)))   # вся старая схема; легенда и врезка остаются
    retext(sl, "Технологическая схема от эталона до районирования; "
               "цвет узла — текущий статус",
           "Одиннадцать блоков технологической схемы; цвет блока — "
           "текущий статус работ")

    CAP_Y = 1.78
    # ---------- стадия 1: подготовка --------------------------------------
    caps(sl, 0.50, CAP_Y, 1.45, "ПОДГОТОВКА", DARK_BOX)
    node(sl, 0.50, 1.98, 1.45, 0.62, "Постановка задачи и критерий оценки",
         num="1", **ST["готово"])
    node(sl, 0.50, 2.76, 1.45, 0.62, "Эталон: воссоздание критериального метода",
         "база сравнения", num="2", **ST["готово"])
    line(sl, 1.225, 2.60, 1.225, 2.76)

    # ---------- стадия 2: данные и знания ---------------------------------
    caps(sl, 2.17, CAP_Y, 2.45, "ДАННЫЕ И ЗНАНИЯ", DARK_BOX)
    node(sl, 2.17, 1.98, 1.10, 0.62, "Агентная сборка базы данных",
         num="3", **ST["в работе"])
    node(sl, 3.52, 1.98, 1.10, 0.62, "Подготовка и паспортизация данных",
         num="4", **ST["в работе"])
    line(sl, 3.31, 2.29, 3.48, 2.29)
    node(sl, 2.17, 2.90, 1.10, 0.62, "Агентная база знаний по публикациям",
         num="5", **ST["в работе"])
    # Граф знаний — «в работе»: мост «база знаний → паспорта» уже работает
    # по черновикам (`scripts/db_kb_evidence.py` в репозитории базы).
    node(sl, 3.52, 2.90, 1.10, 0.62, "Граф знаний и предметные связи",
         num="6", **ST["в работе"])
    line(sl, 3.31, 3.21, 3.48, 3.21)
    # контур возврата 4 -> 3: данных не хватило — возвращаемся к сбору
    line(sl, 4.07, 2.60, 4.07, 2.70, arrow=False)
    line(sl, 4.07, 2.70, 2.72, 2.70, arrow=False)
    line(sl, 2.72, 2.70, 2.72, 2.60)
    tf = textbox(sl, 2.72, 2.71, 1.35, 0.14)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(tf.paragraphs[0], "нужны ещё данные", 7, ACCENT, italic=True)
    # шина «от постановки задачи — обе ветки»
    line(sl, 1.95, 2.29, 2.06, 2.29, arrow=False)
    line(sl, 2.06, 2.29, 2.06, 3.21, arrow=False)
    line(sl, 2.06, 2.29, 2.17, 2.29)
    line(sl, 2.06, 3.21, 2.17, 3.21)

    # ---------- стадия 3: эксперименты ------------------------------------
    caps(sl, 4.84, CAP_Y, 2.30, "ЭКСПЕРИМЕНТЫ", DARK_BOX)
    node(sl, 4.84, 2.05, 2.30, 0.44, "Слияние: выбор направлений обучения",
         num="7", **ST["в работе"])
    line(sl, 5.99, 2.49, 5.99, 2.70)
    # Названия веток — по технологической карте организации
    # (`experiments/build_tech_scheme_presentation.py`, узлы 8A-8E);
    # подписи говорят, чем ветка занята у нас.
    node(sl, 4.84, 2.70, 1.10, 0.46, "Методы из литературы",
         num="8A", **ST["в работе"])
    node(sl, 6.04, 2.70, 1.10, 0.46, "Геопростр. дата-майнинг",
         num="8B", **ST["готово"])
    node(sl, 4.84, 3.20, 1.10, 0.46, "Призн. простр. — инструмент",
         num="8C", **ST["готово"])
    node(sl, 6.04, 3.20, 1.10, 0.46, "Призн. простр. — объяснитель",
         num="8D", **ST["в работе"])
    node(sl, 4.84, 3.70, 2.30, 0.38, "Обучение без учителя + предметный граф",
         num="8E", **ST["в работе"])
    # шины от стадии 2 к блоку 7 и от веток к стадии 4
    line(sl, 4.62, 2.29, 4.73, 2.29, arrow=False)
    line(sl, 4.62, 3.21, 4.73, 3.21, arrow=False)
    line(sl, 4.73, 2.29, 4.73, 3.21, arrow=False)
    line(sl, 4.73, 2.29, 4.84, 2.29)
    line(sl, 7.14, 2.91, 7.25, 2.91, arrow=False)
    line(sl, 7.14, 3.39, 7.25, 3.39, arrow=False)
    line(sl, 7.14, 3.85, 7.25, 3.85, arrow=False)
    line(sl, 7.25, 2.26, 7.25, 3.85, arrow=False)
    line(sl, 7.25, 2.26, 7.36, 2.26)

    # ---------- стадия 4: верификация и итог ------------------------------
    caps(sl, 7.36, CAP_Y, 2.14, "ВЕРИФИКАЦИЯ И ИТОГ", DARK_BOX)
    node(sl, 7.36, 1.98, 2.14, 0.56, "Выбор лучшей модели в каждой ветке",
         num="9", **ST["в работе"])
    node(sl, 7.36, 2.74, 2.14, 0.56,
         "Верификация и соотнесение с предметной областью",
         num="10", **ST["впереди"])
    node(sl, 7.36, 3.50, 2.14, 0.56,
         "Итог: переносимый метод и карта районирования",
         num="11", **ST["впереди"])
    line(sl, 8.43, 2.54, 8.43, 2.74)
    line(sl, 8.43, 3.30, 8.43, 3.50)


# ===========================================================================
def main():
    prs = Presentation(str(SRC))
    sl2, sl3, sl4, sl5 = (prs.slides[i] for i in (1, 2, 3, 4))

    slide2(sl2)
    slide3(sl3)

    # Слайд 4: официальное название раздела.
    retext(sl4, "Где мы сейчас", "Текущие результаты работ")
    for _ in range(2):               # метка в шапке и синяя рубрика
        retext(sl4, "ЭТАП", "РЕЗУЛЬТАТЫ")

    # Слайд 5: официальное название раздела (метки «ПЛАН» и «ДОРОЖНАЯ КАРТА»
    # согласуются с заголовком и остаются как есть).
    retext(sl5, "Что дальше", "План дальнейших работ")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(OUT))
    except PermissionError:
        raise SystemExit(
            f"Файл занят (открыт в PowerPoint): {OUT}\n"
            "Закройте его и запустите скрипт заново — под другим именем "
            "сохранять нельзя.")
    print(f"записано: {OUT}")


if __name__ == "__main__":
    main()
