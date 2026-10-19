# -*- coding: utf-8 -*-
"""Вставка слайдов чужой презентации в колоду проекта с масштабированием холста.

ЗАЧЕМ. Раздел «База знаний и граф» большой колоды сделан отдельно на холсте
13.333 x 7.5", колода серии — 10 x 5.625" (та же пропорция 16:9). Слайды
вставляются без правки содержания: все координаты, размеры, кегли и толщины
линий умножаются на один коэффициент (0.75), вид остаётся прежним.

Поддерживаются текстовые поля, автофигуры, линии (то есть всё, что есть в
исходной презентации); картинки, таблицы и группы НЕ копируются — при встрече
с ними модуль останавливается с понятной ошибкой, чтобы не потерять объект
молча. Заметки докладчика переносятся текстом.

    from src.slides_merge import append_scaled
    n = append_scaled(deck, "файл.pptx")      # deck — src.slides_inst.Deck
"""
import copy

from lxml import etree
from pptx import Presentation

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}

# Атрибуты в EMU / сотых долях пункта, которые масштабируются вместе с холстом.
_EMU_ATTRS = {
    "off": ("x", "y"), "ext": ("cx", "cy"), "chOff": ("x", "y"), "chExt": ("cx", "cy"),
}
_BODY_INSETS = ("lIns", "tIns", "rIns", "bIns")
_SZ_TAGS = ("rPr", "defRPr", "endParaRPr")


def _scale(elm, k, dx=0, dy=0):
    for tag, attrs in _EMU_ATTRS.items():
        for e in elm.iter("{%s}%s" % (A, tag)):
            for a in attrs:
                if e.get(a) is not None:
                    e.set(a, str(int(round(int(e.get(a)) * k))))
    if dx or dy:
        for e in elm.iter("{%s}off" % A):
            e.set("x", str(int(e.get("x")) + int(dx)))
            e.set("y", str(int(e.get("y")) + int(dy)))
    for tag in _SZ_TAGS:
        for e in elm.iter("{%s}%s" % (A, tag)):
            if e.get("sz") is not None:
                e.set("sz", str(max(100, int(round(int(e.get("sz")) * k)))))
    for e in elm.iter("{%s}ln" % A):
        if e.get("w") is not None:
            e.set("w", str(int(round(int(e.get("w")) * k))))
    for e in elm.iter("{%s}bodyPr" % A):
        for a in _BODY_INSETS:
            if e.get(a) is not None:
                e.set(a, str(int(round(int(e.get(a)) * k))))
    for e in elm.iter("{%s}spcPts" % A):
        e.set("val", str(int(round(int(e.get("val")) * k))))
    # межзнаковый интервал rPr/@spc тоже в сотых долях пункта
    for e in elm.iter("{%s}rPr" % A):
        if e.get("spc") is not None:
            e.set("spc", str(int(round(int(e.get("spc")) * k))))


def _recolor(elm, palette):
    """Меняет srgbClr по словарю {HEX: HEX}; цвета вне словаря не трогает."""
    for c in elm.iter("{%s}srgbClr" % A):
        v = c.get("val", "").upper()
        if v in palette:
            c.set("val", palette[v])


def append_scaled(deck, src_path, k=0.75, slides=None, palette=None,
                  line_pt=None, frame=False, text_swap=None, shift=(0.0, 0.0)):
    """Добавляет слайды из `src_path` в конец `deck.prs`; возвращает их число.

    `slides` — номера слайдов исходника с единицы (по умолчанию все).
    `palette` — перекраска {HEX: HEX} (единая цветовая схема колоды);
    `line_pt` — потолок толщины линий, пт; `frame` — поставить колонтитулы
    колоды (верхняя полоса, подвал, номер); `text_swap` — замены слов {старое: новое}
    в текстах (только прогоны, оформление не теряется); `shift` — сдвиг (x, y), дюймы,
    после масштабирования.
    Счётчик страниц `deck.page` увеличивается, чтобы нумерация колоды сошлась.
    """
    src = Presentation(str(src_path))
    layout = deck._layout
    picked = list(range(1, len(src.slides) + 1)) if slides is None else list(slides)
    added = 0
    for num in picked:
        s = src.slides[num - 1]
        for sh in s.shapes:
            if sh.shape_type not in (1, 9, 17):  # автофигура, линия, текстовое поле
                raise SystemExit(
                    "слайд %d: объект %r (%s) не поддерживается вставкой" %
                    (num, sh.name, sh.shape_type))
        new = deck._blank() if frame else deck.prs.slides.add_slide(layout)
        # фон слайда
        bg = s._element.find("p:cSld/p:bg", NS)
        if bg is not None:
            cs = new._element.find("p:cSld", NS)
            old = cs.find("p:bg", NS)
            if old is not None:
                cs.remove(old)
            nb = copy.deepcopy(bg)
            if palette:
                _recolor(nb, palette)
            cs.insert(0, nb)
        # фигуры: при frame оставляем колонтитулы колоды, иначе чистим заготовки макета
        tree = new.shapes._spTree
        if not frame:
            for sh in list(new.shapes):
                tree.remove(sh._element)
        dark = bg is not None and any(
            c.get("val", "").upper() in (palette or {}) for c in bg.iter("{%s}srgbClr" % A)
        ) and False
        for sh in s.shapes:
            el = copy.deepcopy(sh._element)
            _scale(el, k, shift[0] * 914400, shift[1] * 914400)
            if palette:
                _recolor(el, palette)
            if line_pt:
                for e in el.iter("{%s}ln" % A):
                    if e.get("w") is not None:
                        e.set("w", str(min(int(e.get("w")), int(line_pt * 12700))))
            if text_swap:
                for t in el.iter("{%s}t" % A):
                    for old_t, new_t in text_swap.items():
                        if old_t in (t.text or ""):
                            t.text = t.text.replace(old_t, new_t)
            tree.append(el)
        if s.has_notes_slide:
            txt = s.notes_slide.notes_text_frame.text
            if txt.strip():
                new.notes_slide.notes_text_frame.text = txt
        if not frame:
            deck.page += 1  # при frame слайд уже учтён в deck._blank()
        added += 1
    return added
