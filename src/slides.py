# -*- coding: utf-8 -*-
"""Домашний стиль презентаций проекта: одна точка вместо копий в каждом скрипте.

До этого модуля каждый `experiments/build_*presentation*.py` заново определял
одни и те же функции (шапка, текст, пункты, таблица, картинка). Здесь они
собраны в том виде, в каком устоялись к сборке «Обучение модели»
(`build_training_presentation.py`, 09.09.2026): тёмная шапка, золотая линейка,
Calibri, холст 13.333 x 7.5 дюйма, метка блока справа в шапке, номер слайда в
правом углу. Плюс соединители (`build_tech_scheme_presentation.py`) и
редактируемые формулы Office (`build_feature_space_presentation.py`).

Функции оформления принимают слайд первым аргументом и координаты в единицах
python-pptx (`Inches(...)`); нумерацию и шапку ведёт `Deck`.

    from src.slides import Deck, prose, picture, note, Inches
    deck = Deck({1: "ПОСТАНОВКА", 2: "ИТОГИ"})
    deck.title_slide("Название", "Прогноз ...", ["подзаголовок"], "9 сентября 2026")
    s, top = deck.slide("Заголовок", 1)
    prose(s, Inches(0.55), top, Inches(6.3), Inches(5.2), ["абзац"])
    deck.save(OUT)

Проверка готового файла — `src/slide_check.py`.
"""
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DECK_DIR = ROOT / "docs II presentation"
IMG = ROOT / "outputs"

FONT = "Calibri"
GOLD = RGBColor(0xC8, 0x9B, 0x2C)
DARK = RGBColor(0x2B, 0x2B, 0x2B)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF7, 0xF2, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB0, 0x3A, 0x2E)
GREEN = RGBColor(0x2F, 0x6B, 0x4F)
BLUE = RGBColor(0x1F, 0x52, 0x94)
PALE = RGBColor(0xEA, 0xEA, 0xE9)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.55)
CONTENT_W = Inches(12.2)
HEADER_H = Inches(1.0)
NOTE_Y = Inches(6.45)  # типовая верхняя граница врезки-вывода внизу слайда

__all__ = [
    "Deck", "rect", "round_rect", "oval", "connector", "textbox", "set_run",
    "prose", "bullets", "table", "picture", "note", "equation",
    "m_text", "m_plain", "m_sub", "m_sup", "m_frac", "m_rad", "m_paren",
    "Inches", "Pt", "Emu", "RGBColor", "MSO_ANCHOR", "PP_ALIGN",
    "GOLD", "DARK", "GREY", "LIGHT", "WHITE", "RED", "GREEN", "BLUE", "PALE",
    "MARGIN", "CONTENT_W", "NOTE_Y", "IMG", "DECK_DIR", "ROOT",
]


class Deck:
    """Колода в домашнем стиле: холст 16:9, шапки, сквозная нумерация.

    `blocks` — номер блока -> метка в шапке справа («ПОСТАНОВКА», «ИТОГИ»...).
    Титульный слайд не нумеруется, первый слайд с шапкой получает номер 1.
    """

    def __init__(self, blocks=None):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.blocks = blocks or {}
        self._layout = self.prs.slide_layouts[6]
        self._num = 0

    @property
    def sw(self):
        return self.prs.slide_width

    @property
    def sh(self):
        return self.prs.slide_height

    def blank(self):
        """Пустой слайд без шапки и без номера."""
        return self.prs.slides.add_slide(self._layout)

    def title_slide(self, title, lead, lines=(), date=None):
        """Титул: тёмный фон, золотая линейка, заголовок над ней.

        `lead` — первая строка под линейкой (золотом), `lines` — остальные
        (белым), `date` — последняя строка, мельче.
        """
        s = self.blank()
        rect(s, 0, 0, self.sw, self.sh, DARK)
        rect(s, 0, Inches(3.42), self.sw, Pt(5), GOLD)
        tf = textbox(s, Inches(1.0), Inches(1.9), Inches(11.3), Inches(1.5),
                     MSO_ANCHOR.BOTTOM)
        set_run(tf.paragraphs[0].add_run(), title, 34, WHITE, bold=True)
        tf = textbox(s, Inches(1.0), Inches(3.75), Inches(11.3), Inches(2.2))
        rows = [(lead, 17, GOLD, True)] + [(t, 17, WHITE, False) for t in lines]
        if date:
            rows.append((date, 14, WHITE, False))
        for i, (text, size, color, bold) in enumerate(rows):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(14)
            set_run(p.add_run(), text, size, color, bold=bold)
        return s

    def slide(self, title, block=None, subtitle=None):
        """Слайд с шапкой; возвращает слайд и верхнюю границу контента."""
        s = self.blank()
        self._num += 1
        rect(s, 0, 0, self.sw, HEADER_H, DARK)
        rect(s, 0, HEADER_H, self.sw, Pt(4), GOLD)
        tf = textbox(s, MARGIN, 0, Inches(9.3), HEADER_H, MSO_ANCHOR.MIDDLE)
        set_run(tf.paragraphs[0].add_run(), title, 23, WHITE, bold=True)
        if block is not None:
            tfc = textbox(s, Inches(9.6), 0, Inches(2.7), HEADER_H,
                          MSO_ANCHOR.MIDDLE)
            tfc.paragraphs[0].alignment = PP_ALIGN.RIGHT
            set_run(tfc.paragraphs[0].add_run(), self.blocks[block], 11, GOLD,
                    bold=True)
        tfn = textbox(s, Inches(12.35), 0, Inches(0.75), HEADER_H,
                      MSO_ANCHOR.MIDDLE)
        tfn.paragraphs[0].alignment = PP_ALIGN.RIGHT
        set_run(tfn.paragraphs[0].add_run(), str(self._num), 18, GOLD, bold=True)
        top = Inches(1.35)
        if subtitle:
            tfs = textbox(s, MARGIN, Inches(1.28), CONTENT_W, Inches(0.5))
            set_run(tfs.paragraphs[0].add_run(), subtitle, 14, GREY, italic=True)
            top = Inches(1.9)
        return s, top

    def save(self, path):
        path = Path(path)
        self.prs.save(str(path))
        try:
            shown = path.relative_to(ROOT)
        except ValueError:
            shown = path
        print("OK: %s, слайдов %d" % (shown, len(self.prs.slides._sldIdLst)))
        return path


# ------------------------------------------------------------------ фигуры
def rect(slide, x, y, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def round_rect(slide, x, y, w, h, color, radius=0.06):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = radius
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def oval(slide, x, y, d, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def connector(slide, x1, y1, x2, y2, color, weight=1.75, arrow=True):
    """Ломаная линия-соединитель, по умолчанию со стрелкой на конце."""
    cxn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW, x1, y1, x2, y2)
    cxn.line.color.rgb = color
    cxn.line.width = Pt(weight)
    if arrow:
        ln = cxn.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    return cxn


# ------------------------------------------------------------------- текст
def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def set_run(r, text, size, color=DARK, bold=False, italic=False):
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = FONT


def prose(s, x, y, w, h, paragraphs, size=16):
    """Сплошной текст абзацами, без маркеров."""
    tf = textbox(s, x, y, w, h)
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(13)
        p.line_spacing = 1.15
        set_run(p.add_run(), text, size, DARK)
    return tf


def bullets(s, x, y, w, h, items, size=15.5):
    """Пункты с золотым маркером; элемент — строка или пара (жирное, хвост)."""
    tf = textbox(s, x, y, w, h)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(11)
        p.line_spacing = 1.1
        set_run(p.add_run(), "— ", size, GOLD, bold=True)
        if isinstance(item, tuple):
            set_run(p.add_run(), item[0], size, DARK, bold=True)
            set_run(p.add_run(), item[1], size, DARK)
        else:
            set_run(p.add_run(), item, size, DARK)
    return tf


def table(s, x, y, w, rows, col_w=None, size=13, head_size=12.5, row_h=0.42,
          accent=None):
    """Таблица без сетки PowerPoint: тёмная шапка, чередование фона строк.

    `rows[0]` — шапка; `accent` — номер строки данных (с нуля), которую
    выделить золотым фоном. Возвращает нижнюю границу таблицы.
    """
    n_col = len(rows[0])
    col_w = col_w or [w / n_col] * n_col
    head_h = Inches(row_h + 0.08)
    rect(s, x, y, w, head_h, DARK)
    cx = x
    for j, cell in enumerate(rows[0]):
        tf = textbox(s, cx + Inches(0.08), y, col_w[j] - Inches(0.16), head_h,
                     MSO_ANCHOR.MIDDLE)
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        set_run(tf.paragraphs[0].add_run(), cell, head_size, WHITE, bold=True)
        cx += col_w[j]
    ry = y + head_h
    for i, row in enumerate(rows[1:]):
        h = Inches(row_h)
        bg = LIGHT if accent == i else (PALE if i % 2 == 0 else WHITE)
        rect(s, x, ry, w, h, bg)
        if accent == i:
            rect(s, x, ry, Pt(3.5), h, GOLD)
        cx = x
        for j, cell in enumerate(row):
            bold = accent == i and j > 0
            tf = textbox(s, cx + Inches(0.08), ry, col_w[j] - Inches(0.16), h,
                         MSO_ANCHOR.MIDDLE)
            tf.paragraphs[0].alignment = (PP_ALIGN.LEFT if j == 0
                                          else PP_ALIGN.CENTER)
            set_run(tf.paragraphs[0].add_run(), cell, size, DARK, bold=bold)
            cx += col_w[j]
        ry += h
    return ry


def picture(s, name, x, y, max_w, max_h, caption=None):
    """Картинка вписывается в рамку с сохранением пропорций, по центру.

    `name` — путь относительно `outputs/` или абсолютный. Возвращает нижнюю
    границу картинки (без подписи).
    """
    path = Path(name)
    if not path.is_absolute():
        path = IMG / path
    iw, ih = Image.open(path).size
    k = min(max_w / iw, max_h / ih)
    w, h = Emu(int(iw * k)), Emu(int(ih * k))
    px = x + Emu(int((max_w - int(iw * k)) / 2))
    s.shapes.add_picture(str(path), px, y, width=w, height=h)
    if caption:
        tf = textbox(s, x, y + h + Pt(3), max_w, Inches(0.4))
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        set_run(tf.paragraphs[0].add_run(), caption, 10.5, GREY, italic=True)
    return y + h


def note(s, x, y, w, text, color=GOLD, size=13.5, h=None):
    """Врезка-вывод: цветная полоса слева + жирный текст на светлом фоне.

    Высота по умолчанию прикидывается по длине текста (для ширины ~12 дюймов
    и кегля 13.5); при другом размере лучше задать `h` явно и проверить
    `src/slide_check.py`.
    """
    if h is None:
        h = Inches(0.42 + 0.26 * (len(text) // 105))
    round_rect(s, x, y, w, h, LIGHT, radius=0.05)
    rect(s, x, y, Pt(4), h, color)
    tf = textbox(s, x + Inches(0.22), y, w - Inches(0.4), h, MSO_ANCHOR.MIDDLE)
    set_run(tf.paragraphs[0].add_run(), text, size, DARK, bold=True)
    return y + h


# ----------------------------------------------------------------- формулы
# Редактируемые формулы Office (OMML). Выражение собирается из кусков m_*:
#   m_text("ρ = 1 − ") + m_frac(m_text("6Σ") + m_sup(m_sub("d", "i"), "2"),
#                               m_text("n(n²−1)"))
# и вставляется отдельным абзацем: equation(tf, omml, "ρ = 1 − 6Σd²/...", 18).
_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"


def m_text(text):
    """Обычный (курсивный) фрагмент формулы — переменные, знаки."""
    return '<m:r><m:t xml:space="preserve">%s</m:t></m:r>' % text


def m_plain(text):
    """Прямой фрагмент — имена функций (cos, sin, arctan2)."""
    return ('<m:r><m:rPr><m:sty m:val="p"/></m:rPr>'
            '<m:t xml:space="preserve">%s</m:t></m:r>' % text)


def m_sub(base, sub):
    return ('<m:sSub><m:e>%s</m:e><m:sub>%s</m:sub></m:sSub>'
            % (m_text(base), m_text(sub)))


def m_sup(base_xml, sup_text):
    return ('<m:sSup><m:e>%s</m:e><m:sup>%s</m:sup></m:sSup>'
            % (base_xml, m_text(sup_text)))


def m_frac(num_xml, den_xml):
    return ('<m:f><m:fPr><m:type m:val="bar"/></m:fPr>'
            '<m:num>%s</m:num><m:den>%s</m:den></m:f>' % (num_xml, den_xml))


def m_rad(inner_xml):
    """Квадратный корень."""
    return ('<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/>'
            '<m:e>%s</m:e></m:rad>' % inner_xml)


def m_paren(inner_xml, left="(", right=")"):
    return ('<m:d><m:dPr><m:begChr m:val="%s"/><m:endChr m:val="%s"/></m:dPr>'
            '<m:e>%s</m:e></m:d>' % (left, right, inner_xml))


def equation(tf, omml, fallback, size, first=False, color="C89B2C",
             space_after=8):
    """Абзац с формулой Office; `fallback` — текст для программ без OMML."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    xml = (
        f'<mc:AlternateContent xmlns:mc="{_MC_NS}">'
        f'<mc:Choice xmlns:a14="{_A14_NS}" Requires="a14"><a14:m>'
        f'<m:oMathPara xmlns:m="{_M_NS}"><m:oMath>{omml}</m:oMath></m:oMathPara>'
        f'</a14:m></mc:Choice>'
        f'<mc:Fallback><a:r xmlns:a="{_A_NS}">'
        f'<a:rPr lang="ru-RU" b="1" sz="{int(round(size * 100))}">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{FONT}"/></a:rPr>'
        f'<a:t>{fallback}</a:t></a:r></mc:Fallback>'
        f'</mc:AlternateContent>'
    )
    p._p.append(etree.fromstring(xml.encode("utf-8")))
    return p
